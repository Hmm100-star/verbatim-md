"""Deterministic DOCX rendering without quotation transformations or macros."""
from __future__ import annotations

from bisect import bisect_right
from datetime import datetime, timezone
from pathlib import Path
import os
import tempfile

from docx import Document as WordDocument
from docx.oxml import parse_xml
from docx.shared import Inches, Pt
from defusedxml import ElementTree as ET

from .config import Profile, load_profile
from .evidence import attach_manifest, read_package
from .schema import Analytic, Card, Document, Heading, Span
from .styles import CITE_STYLE, UNDERLINE_STYLE, ensure_style_contract, setup_styles

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _import_styles_unchecked(word, template: str | Path) -> None:
    """Read only compatible style definitions; never import contents or macros."""
    if Path(template).suffix.lower() != '.docx':
        raise ValueError('reference template must be a macro-free .docx')
    parts = read_package(template)
    if any('vba' in name.lower() for name in parts):
        raise ValueError('reference template contains macro parts')
    root = ET.fromstring(parts['word/styles.xml'], forbid_dtd=True)
    allowed = {'Normal', 'Heading1', 'Heading2', 'Heading3', 'Heading4',
               CITE_STYLE, UNDERLINE_STYLE, 'Emphasis'}
    seen = set()
    known = {node.get(W + 'styleId') for node in word.styles.element} | allowed
    for node in root:
        identifier = node.get(W + 'styleId')
        if identifier not in allowed:
            continue
        if identifier in seen:
            raise ValueError(f'duplicate reference style {identifier}')
        seen.add(identifier)
        for relation in ('basedOn', 'next', 'link'):
            reference = node.find(W + relation)
            if reference is not None and reference.get(W + 'val') not in known:
                raise ValueError(f'reference style {identifier} depends on an unsupported style')
        expected_name = {
            'Normal': 'normal', 'Heading1': 'heading 1', 'Heading2': 'heading 2',
            'Heading3': 'heading 3', 'Heading4': 'heading 4',
            CITE_STYLE: 'style 13 pt bold', UNDERLINE_STYLE: 'style underline',
            'Emphasis': 'emphasis',
        }[identifier]
        name = node.find(W + 'name')
        expected_type = 'paragraph' if identifier in {'Normal', 'Heading1', 'Heading2', 'Heading3', 'Heading4'} else 'character'
        if (name is None or name.get(W + 'val', '').casefold() != expected_name
                or node.get(W + 'type') != expected_type):
            raise ValueError(f'reference style {identifier} has an incompatible name or type')
        # Styles are passive XML only: no external relationship references.
        if any('relationships' in key for child in node.iter() for key in child.attrib):
            raise ValueError('reference styles contain relationship references')
        replacement = parse_xml(ET.tostring(node, encoding='utf-8'))
        for old in list(word.styles.element):
            if old.get(W + 'styleId') == identifier:
                word.styles.element.remove(old)
        word.styles.element.append(replacement)



def _import_styles(word, template: str | Path) -> None:
    from zipfile import BadZipFile
    from xml.etree.ElementTree import ParseError
    from defusedxml.common import DefusedXmlException
    try:
        _import_styles_unchecked(word, template)
    except (BadZipFile, KeyError, ParseError, DefusedXmlException, UnicodeError) as exc:
        raise ValueError(f'Invalid reference DOCX: {exc}') from exc


def _formatting(spans: tuple[Span, ...]) -> tuple[list[int], list[frozenset[str]]]:
    """Sweep endpoints once; counters preserve overlaps of the same attribute."""
    events: dict[int, list[tuple[str, int]]] = {0: []}
    for span in spans:
        events.setdefault(span.start, []).append((span.kind, 1))
        events.setdefault(span.end, []).append((span.kind, -1))
    points = sorted(events)
    active: dict[str, int] = {}
    states = []
    for point in points:
        for kind, change in events[point]:
            active[kind] = active.get(kind, 0) + change
        states.append(frozenset(kind for kind, count in active.items() if count > 0))
    return points, states


def _runs(paragraph, text: str, spans: tuple[Span, ...], offset: int,
          profile: Profile, *, evidence: bool = False,
          formatting: tuple[list[int], list[frozenset[str]]] | None = None) -> None:
    end = offset + len(text)
    points, states = formatting if formatting is not None else _formatting(spans)
    index = bisect_right(points, offset) - 1
    start = offset
    while start < end:
        stop = min(end, points[index + 1]) if index + 1 < len(points) else end
        kinds = states[index]
        run = paragraph.add_run(text[start - offset:stop - offset])
        if 'cite' in kinds:
            run.style = 'Style 13 pt Bold'
        elif 'emphasis' in kinds:
            run.style = 'Emphasis'
        elif 'underline' in kinds:
            run.style = 'Style Underline'
        if 'underline' in kinds or 'emphasis' in kinds:
            run.font.underline = True
        if 'emphasis' in kinds:
            run.font.bold = True
        if 'highlight' in kinds:
            from docx.oxml import OxmlElement
            highlight = OxmlElement('w:highlight')
            highlight.set(W + 'val', profile.highlight_color)
            run._r.get_or_add_rPr().append(highlight)
        # Mirror documented Verbatim Shrink selection: non-underlined text.
        if 'small' in kinds or (evidence and profile.shrink_unread and not kinds.intersection({'underline', 'emphasis'})):
            run.font.size = Pt(profile.small_size)
        elif evidence:
            run.font.size = Pt(profile.font_size)
        start = stop
        index += 1


def render_document(document: Document, output: str | Path,
                    profile: Profile | str | Path | None = None,
                    template: str | Path | None = None) -> Path:
    """Render validated immutable input and embed preservation metadata.

    Output replacement is atomic. Overwrite authorization belongs to API/CLI.
    A reference DOCX imports only eight recognized style definitions.
    """
    profile = load_profile(profile)
    from .validator import validate_document
    report = validate_document(document)
    if not report['valid']:
        raise ValueError('; '.join(report['errors']))
    word = WordDocument()
    setup_styles(word, profile)
    if template is not None:
        _import_styles(word, template)
    ensure_style_contract(word)
    for section in word.sections:
        section.top_margin = section.bottom_margin = Inches(profile.margin_inches)
        section.left_margin = section.right_margin = Inches(profile.margin_inches)
    word.core_properties.title = document.title
    word.core_properties.author = ''
    word.core_properties.last_modified_by = ''
    fixed_date = datetime(2000, 1, 1, tzinfo=timezone.utc)
    word.core_properties.created = word.core_properties.modified = fixed_date
    sources = {source.id: source for source in document.sources}
    card_paragraphs: dict[str, list[int]] = {}
    paragraph_index = 0
    for block in document.blocks:
        if isinstance(block, Heading):
            word.add_paragraph(block.text, style=f'Heading {block.level}')
            paragraph_index += 1
        elif isinstance(block, Analytic):
            for text in block.text.split('\n'):
                word.add_paragraph(text, style='Normal')
                paragraph_index += 1
        elif isinstance(block, Card):
            source = sources[block.source_id]
            citation = word.add_paragraph(style='Normal')
            paragraph_index += 1
            citation.paragraph_format.keep_with_next = True
            _runs(citation, source.citation, source.citation_spans, 0, profile)
            indices = []
            offset = 0
            formatting = _formatting(block.spans)
            for text in block.text.split('\n'):
                indices.append(paragraph_index)
                paragraph_index += 1
                paragraph = word.add_paragraph(style='Normal')
                _runs(paragraph, text, block.spans, offset, profile, evidence=True,
                      formatting=formatting)
                offset += len(text) + 1
            card_paragraphs[block.id] = indices
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(suffix='.docx', dir=path.parent)
    os.close(descriptor)
    try:
        word.save(temporary)
        attach_manifest(temporary, document, card_paragraphs)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return path
